"""PDF merging with Unicode TOC pages and matching bookmarks."""
import math
import os
import difflib

import fitz
import pandas as pd


def metadata_entries(metadata_path, folder, output_path):
    if metadata_path.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        workbook = load_workbook(metadata_path, read_only=True, data_only=True)
        try:
            rows = iter(workbook.active.iter_rows(values_only=True))
            headers = [str(value or '').strip() for value in next(rows, ())]
            df = pd.DataFrame(list(rows), columns=headers).fillna('')
        finally:
            workbook.close()
    else:
        df = pd.read_csv(metadata_path, dtype=str, encoding='utf-8-sig').fillna('')
    df.columns = df.columns.str.strip()
    canonical = {key.lower(): key for key in ('Order', 'TFL', 'FileName', 'Title')}
    df = df.rename(columns={key: canonical.get(key.lower(), key) for key in df.columns})
    required = {'Order', 'TFL', 'FileName', 'Title'}
    if not required.issubset(df.columns):
        raise ValueError('Metadata requires: Order, TFL, FileName, Title')
    entries, orders = [], set()
    available = {}
    for item in os.listdir(folder):
        if os.path.isfile(os.path.join(folder, item)) and item.lower().endswith(('.rtf', '.pdf')):
            available.setdefault(item.upper(), []).append(item)
    for index, row in df.iterrows():
        if not any(str(row[c]).strip() for c in required):
            continue
        label = f'Metadata row {index + 2}'
        try:
            order = float(row['Order'])
        except ValueError:
            raise ValueError(f'{label}: Order must be numeric') from None
        if not math.isfinite(order) or order in orders:
            raise ValueError(f'{label}: Order must be finite and unique')
        orders.add(order)
        if str(row['TFL']).strip().upper() not in ('T', 'F', 'L'):
            raise ValueError(f'{label}: TFL must be T, F or L')
        name, title = str(row['FileName']).strip(), str(row['Title']).strip()
        if not name or name in ('.', '..') or any(char in name for char in ('/', '\\', ':')):
            raise ValueError(f'{label}: FileName must be a filename without a folder')
        if not title:
            raise ValueError(f'{label}: Title is required')
        # An extensionless name selects the original RTF first, otherwise a PDF.
        candidates = [name] if name.lower().endswith(('.pdf', '.rtf')) else [name + '.rtf', name + '.pdf']
        matches = next((available[item.upper()] for item in candidates if item.upper() in available), [])
        if not matches:
            near = difflib.get_close_matches(candidates[0].upper(), available, n=3, cutoff=0.5)
            suggestions = ', '.join(available[key][0] for key in near)
            raise ValueError(f'{label}: RTF/PDF not found: {name}\nFolder: {folder}\n'
                             'Matching ignores letter case. Check spelling and separators.'
                             + (f'\nSimilar filenames: {suggestions}' if suggestions else ''))
        if len(matches) > 1:
            raise ValueError(f'{label}: Ambiguous filenames after uppercase matching: {matches}')
        match = matches[0]
        path = os.path.join(folder, match)
        if os.path.normcase(os.path.abspath(path)) == os.path.normcase(os.path.abspath(output_path)):
            raise ValueError(f'{label}: input filename equals output filename')
        entries.append((order, path, title))
    if not entries:
        raise ValueError('Metadata has no input entries')
    return [(path, title) for _, path, title in sorted(entries)]


def merge_with_toc(entries, output_path, password=None):
    if not entries:
        raise ValueError('No PDFs to combine')
    font = fitz.Font('china-s')
    layout, toc_page, y = [], 0, 90
    # Wrap by rendered width, including Chinese text, before determining page offsets.
    for _, title in entries:
        lines, line = [], ''
        for char in title.replace('\r', ' ').replace('\n', ' '):
            if line and font.text_length(line + char, fontsize=11) > 440:
                lines.append(line)
                line = ''
            line += char
        lines.append(line)
        for number, line in enumerate(lines):
            if y > 770:
                toc_page, y = toc_page + 1, 90
            layout.append((toc_page, y, line, number == 0))
            y += 18
        y += 9
    toc_count = toc_page + 1
    with fitz.open() as result:
        for _ in range(toc_count):
            page = result.new_page(width=595, height=842)
            page.insert_font(fontname='tocfont', fontbuffer=font.buffer)
            page.insert_text((48, 48), 'Table of Contents', fontsize=20)
        targets, bookmarks = [], [[1, 'Table of Contents', 1]]
        for path, title in entries:
            with fitz.open(path) as source:
                if not len(source):
                    raise ValueError(f'Empty PDF: {os.path.basename(path)}')
                target = len(result)
                targets.append(target)
                result.insert_pdf(source)
                bookmarks.append([1, title, target + 1])
        entry_index = -1
        for page_index, baseline, text, first_line in layout:
            if first_line:
                entry_index += 1
            page = result[page_index]
            page.insert_text((48, baseline), text, fontsize=11, fontname='tocfont')
            if first_line:
                page.insert_text((525, baseline), str(targets[entry_index] + 1), fontsize=11)
            page.insert_link({'kind': fitz.LINK_GOTO,
                              'from': fitz.Rect(46, baseline - 13, 555, baseline + 4),
                              'page': targets[entry_index], 'to': fitz.Point(0, 0)})
        result.set_toc(bookmarks)
        options = dict(garbage=4, deflate=True)
        if password:
            options.update(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=password)
        result.save(output_path, **options)
    return toc_count
